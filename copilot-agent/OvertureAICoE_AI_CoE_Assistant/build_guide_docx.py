"""Render AI_CoE_Assistant_Guide.md to AI_CoE_Assistant_Guide.docx (the file uploaded as the agent's knowledge).

    py -3 build_guide_docx.py

Handles the guide's subset of Markdown: # and ## headings, paragraphs, "- " bullets, "1. " numbered items
and their indented "- " sub-bullets. Edit the .md, then rerun; never edit the .docx by hand.
"""
import re
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'AI_CoE_Assistant_Guide.md'
TARGET = HERE / 'AI_CoE_Assistant_Guide.docx'


def main():
    doc = Document()
    doc.core_properties.title = 'AI CoE Assistant Guide'
    for line in SOURCE.read_text(encoding='utf-8').splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('   - '):
            doc.add_paragraph(line[5:].strip(), style='List Bullet 2')
        elif line.startswith('   '):
            doc.add_paragraph(line.strip())
        elif line.startswith('- '):
            doc.add_paragraph(line[2:].strip(), style='List Bullet')
        elif re.match(r'\d+\. ', line):
            doc.add_paragraph(re.sub(r'^\d+\. ', '', line).strip(), style='List Number')
        else:
            doc.add_paragraph(line.strip())
    doc.save(TARGET)
    print('Wrote', TARGET)


if __name__ == '__main__':
    main()
